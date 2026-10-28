from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.constants import CATEGORY_NAMES
from app.database.repository import (
    get_budget_expenses_for_export,
    get_budget_incomes_for_export,
)

EXPORT_DIR = Path("/opt/expense-bot/exports")

INCOME_NAMES = {
    "ilya": "Илья",
    "liza": "Лиза",
}


def fill_sheet(sheet, rows):
    headers = [
        "Дата", "Категория", "Сумма",
        "Кто добавил", "Комментарий",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid", fgColor="244062"
        )
        cell.alignment = Alignment(horizontal="center")

    total = Decimal("0.00")

    for item in rows:
        sheet.append(item)
        total += Decimal(str(item[2]))

    sheet.append([])
    sheet.append(["", "ИТОГО", float(total)])

    total_row = sheet.max_row
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)

    for row in range(2, total_row + 1):
        sheet.cell(row, 3).number_format = '#,##0.00'

    for index, width in enumerate(
        (16, 30, 20, 24, 45), start=1
    ):
        sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{max(1, total_row - 2)}"

    return total


async def create_budget_xlsx(
    budget_id: int,
    date_from: date,
    date_to: date,
) -> Path:
    expenses = await get_budget_expenses_for_export(
        budget_id, date_from, date_to
    )
    incomes = await get_budget_incomes_for_export(
        budget_id, date_from, date_to
    )

    workbook = Workbook()

    expense_sheet = workbook.active
    expense_sheet.title = "Расходы"

    expense_total = fill_sheet(
        expense_sheet,
        [
            [
                item.expense_date.strftime("%d.%m.%Y"),
                CATEGORY_NAMES.get(item.category, item.category),
                float(item.amount),
                item.display_name,
                item.comment or "",
            ]
            for item in expenses
        ],
    )

    income_sheet = workbook.create_sheet("Доходы")

    income_total = fill_sheet(
        income_sheet,
        [
            [
                item.income_date.strftime("%d.%m.%Y"),
                INCOME_NAMES.get(item.category, item.category),
                float(item.amount),
                item.display_name,
                item.comment or "",
            ]
            for item in incomes
        ],
    )

    summary = workbook.create_sheet("Сводка", 0)
    summary.append(["Семейный бюджет", "Сумма, ₽"])
    summary.append([
        "Период",
        f"{date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}",
    ])
    summary.append(["Доходы", float(income_total)])
    summary.append(["Расходы", float(expense_total)])
    summary.append([
        "Баланс",
        float(income_total - expense_total),
    ])

    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")

    summary["A5"].font = Font(bold=True)
    summary["B5"].font = Font(bold=True)

    for row in range(3, 6):
        summary.cell(row, 2).number_format = '#,##0.00'

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 34

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    # Уникальный временный файл: выгрузки разных
    # пользователей не перезапишут друг друга.
    with NamedTemporaryFile(
        dir=EXPORT_DIR,
        prefix=f"budget_{date_from}_{date_to}_",
        suffix=".xlsx",
        delete=False,
    ) as temporary:
        path = Path(temporary.name)

    try:
        workbook.save(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    return path
