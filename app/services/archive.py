from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from app.constants import CATEGORY_NAMES
from app.database.repository import (
    get_budget_expenses_for_export,
    mark_month_archived,
)
from app.services.google_drive import (
    download_archive,
    upload_or_update_archive,
)


ARCHIVE_LOCAL_FILE = Path(
    "/opt/expense-bot/exports/family_expenses.xlsx"
)


MONTH_NAMES = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}


def money_to_excel(
    amount: Decimal,
) -> float:
    return float(amount)


def get_month_bounds(
    year: int,
    month: int,
) -> tuple[date, date]:
    date_from = date(
        year,
        month,
        1,
    )

    if month == 12:
        next_month = date(
            year + 1,
            1,
            1,
        )
    else:
        next_month = date(
            year,
            month + 1,
            1,
        )

    date_to = date.fromordinal(
        next_month.toordinal() - 1
    )

    return date_from, date_to


def create_empty_workbook() -> Workbook:
    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    return workbook


def get_or_create_year_sheet(
    workbook: Workbook,
    year: int,
):
    sheet_name = str(year)

    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]

    sheet = workbook.create_sheet(
        title=sheet_name,
    )

    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 32
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 22

    return sheet


def clear_month_block(
    sheet,
    month_marker: str,
) -> None:
    marker_row = None
    next_marker_row = None

    for row_number in range(
        1,
        sheet.max_row + 1,
    ):
        value = sheet.cell(
            row=row_number,
            column=1,
        ).value

        if value == month_marker:
            marker_row = row_number
            continue

        if (
            marker_row is not None
            and isinstance(value, str)
            and value.startswith("MONTH:")
        ):
            next_marker_row = row_number
            break

    if marker_row is None:
        return

    if next_marker_row is None:
        next_marker_row = (
            sheet.max_row + 1
        )

    sheet.delete_rows(
        marker_row,
        next_marker_row - marker_row,
    )


async def write_month_to_sheet(
    sheet,
    budget_id: int,
    year: int,
    month: int,
) -> None:
    month_name = MONTH_NAMES[month]

    month_marker = (
        f"MONTH:{year}-{month:02d}"
    )

    clear_month_block(
        sheet=sheet,
        month_marker=month_marker,
    )

    date_from, date_to = get_month_bounds(
        year=year,
        month=month,
    )

    expenses = await get_budget_expenses_for_export(
        budget_id=budget_id,
        date_from=date_from,
        date_to=date_to,
    )

    start_row = sheet.max_row + 1

    if start_row > 1:
        start_row += 1

    sheet.cell(
        row=start_row,
        column=1,
        value=month_marker,
    )

    sheet.row_dimensions[start_row].hidden = True

    title_row = start_row + 1

    sheet.cell(
        row=title_row,
        column=1,
        value=f"{month_name.upper()} {year}",
    )

    sheet.cell(
        row=title_row,
        column=1,
    ).font = Font(
        bold=True,
        size=14,
    )

    header_row = title_row + 1

    headers = [
        "Дата",
        "Категория",
        "Сумма",
        "Кто добавил",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=header_row,
            column=column_number,
            value=header,
        )

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    current_row = header_row + 1

    total_amount = Decimal("0.00")
    category_totals: dict[str, Decimal] = {}

    for expense in expenses:
        category_name = CATEGORY_NAMES.get(
            expense.category,
            expense.category,
        )

        sheet.cell(
            row=current_row,
            column=1,
            value=expense.expense_date.strftime(
                "%d.%m.%Y"
            ),
        )

        sheet.cell(
            row=current_row,
            column=2,
            value=category_name,
        )

        sheet.cell(
            row=current_row,
            column=3,
            value=money_to_excel(
                expense.amount
            ),
        )

        sheet.cell(
            row=current_row,
            column=4,
            value=expense.display_name,
        )

        sheet.cell(
            row=current_row,
            column=3,
        ).number_format = '#,##0.00'

        total_amount += expense.amount

        category_totals[expense.category] = (
            category_totals.get(
                expense.category,
                Decimal("0.00"),
            )
            + expense.amount
        )

        current_row += 1

    if not expenses:
        sheet.cell(
            row=current_row,
            column=1,
            value="Расходов за месяц нет.",
        )

        current_row += 1

    current_row += 1

    sheet.cell(
        row=current_row,
        column=2,
        value="ИТОГО ЗА МЕСЯЦ",
    ).font = Font(
        bold=True,
    )

    sheet.cell(
        row=current_row,
        column=3,
        value=money_to_excel(
            total_amount
        ),
    ).font = Font(
        bold=True,
    )

    sheet.cell(
        row=current_row,
        column=3,
    ).number_format = '#,##0.00'

    current_row += 2

    if category_totals:
        sheet.cell(
            row=current_row,
            column=2,
            value="По категориям:",
        ).font = Font(
            bold=True,
        )

        current_row += 1

        sorted_categories = sorted(
            category_totals.items(),
            key=lambda item: item[1],
            reverse=True,
        )

        for category, amount in sorted_categories:
            category_name = CATEGORY_NAMES.get(
                category,
                category,
            )

            sheet.cell(
                row=current_row,
                column=2,
                value=category_name,
            )

            sheet.cell(
                row=current_row,
                column=3,
                value=money_to_excel(
                    amount
                ),
            )

            sheet.cell(
                row=current_row,
                column=3,
            ).number_format = '#,##0.00'

            current_row += 1


async def update_month_archive(
    budget_id: int,
    year: int,
    month: int,
) -> str:
    """Обновить месяц в архиве и отметить его как архивированный."""

    ARCHIVE_LOCAL_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    try:
        downloaded = download_archive(
            destination=ARCHIVE_LOCAL_FILE,
        )

        if downloaded:
            workbook = load_workbook(
                ARCHIVE_LOCAL_FILE
            )
        else:
            workbook = create_empty_workbook()

        sheet = get_or_create_year_sheet(
            workbook=workbook,
            year=year,
        )

        await write_month_to_sheet(
            sheet=sheet,
            budget_id=budget_id,
            year=year,
            month=month,
        )

        workbook.save(
            ARCHIVE_LOCAL_FILE
        )

        file_id = upload_or_update_archive(
            local_file=ARCHIVE_LOCAL_FILE,
        )

        # Только ПОСЛЕ успешной загрузки на Google Drive
        # фиксируем месяц как архивированный.
        await mark_month_archived(
            budget_id=budget_id,
            year=year,
            month=month,
        )

        return file_id

    finally:
        ARCHIVE_LOCAL_FILE.unlink(
            missing_ok=True,
        )
